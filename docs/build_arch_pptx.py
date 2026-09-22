#!/usr/bin/env python3
"""Rebuild docs/superior_arch_cio.html as ONE native 16:9 slide (PPTX) for
Google Slides. Geometry (docs/_arch_geom.json) was extracted from the live DOM,
so every element maps 1:1: the slide is 13.333x7.5in and 1 CSS px = 9144 EMU."""
import json, subprocess, os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
GEOM = json.load(open(os.path.join(HERE, "_arch_geom.json")))
EMU_PX = 9144                      # exact at 13.333in / 1333px
DARK = (28, 61, 80)                # ontology row backdrop for rgba compositing
def E(px): return Emu(int(round(px * EMU_PX)))
def PT(px): return Pt(px * 0.75)   # CSS px -> pt

def parse_color(s, bd=DARK):
    s = s.strip()
    if s.startswith("rgba"):
        r, g, b, a = [float(x) for x in s[s.find("(")+1:s.find(")")].split(",")]
    elif s.startswith("rgb"):
        r, g, b = [float(x) for x in s[s.find("(")+1:s.find(")")].split(",")]; a = 1.0
    else:
        return RGBColor(0, 0, 0)
    r = r*a + bd[0]*(1-a); g = g*a + bd[1]*(1-a); b = b*a + bd[2]*(1-a)
    return RGBColor(int(round(r)), int(round(g)), int(round(b)))

prs = Presentation()
prs.slide_width = Emu(12192000)
prs.slide_height = Emu(6858000)
slide = prs.slides.add_slide(prs.slide_layouts[6])   # blank

def no_shadow(shape):
    el = shape._element.spPr
    if el.find(qn('a:effectLst')) is None:
        el.append(el.makeelement(qn('a:effectLst'), {}))

def add_box(b):
    x, y, w, h, rx = b["x"], b["y"], b["w"], b["h"], b.get("rx", 0)
    bg, bc, bw = b.get("bg"), b.get("bc"), b.get("bw", 0)
    # row separators: bordered, no fill, square, full-width -> draw top hairline only
    if bg is None and bc and rx == 0 and w > 1000:
        ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(max(bw, 1)))
        ln.fill.solid(); ln.fill.fore_color.rgb = parse_color(bc)
        ln.line.fill.background(); no_shadow(ln); return
    # choose geometry
    is_circle = abs(w - h) < 2 and rx >= 25
    if is_circle:
        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(x), E(y), E(w), E(h))
    elif rx > 0.5:
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(x), E(y), E(w), E(h))
        try:
            adj = 0.5 if (rx >= 25 or rx*2 >= min(w, h)) else rx/min(w, h)
            shp.adjustments[0] = min(0.5, adj)
        except Exception:
            pass
    else:
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
    if bg:
        shp.fill.solid(); shp.fill.fore_color.rgb = parse_color(bg)
    else:
        shp.fill.background()
    if bc and bw:
        shp.line.color.rgb = parse_color(bc); shp.line.width = Pt(bw * 0.75)
    else:
        shp.line.fill.background()
    no_shadow(shp)

def add_text_runs(x, y, w, h, runs, align="start", anchor=MSO_ANCHOR.MIDDLE, wrap=True):
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    try: tf.auto_size = None
    except Exception: pass
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = {"right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}.get(align, PP_ALIGN.LEFT)
    p.line_spacing = 1.0
    for text, fs, fw, color, ff, italic in runs:
        r = p.add_run(); r.text = text
        f = r.font
        f.size = PT(fs); f.name = ff; f.italic = italic
        f.bold = int(str(fw)) >= 700
        f.color.rgb = parse_color(color) if isinstance(color, str) else color
    return tb

def T(t):
    txt = t["text"]
    if t.get("tt") == "uppercase":
        txt = txt.upper()
    add_text_runs(t["x"], t["y"], t["w"], t["h"],
                  [(txt, t["fs"], t["fw"], t["color"], t["ff"], t.get("italic", False))],
                  align=t.get("ta", "start"))

# ---- draw boxes (DOM/paint order) ----
for b in GEOM["boxes"]:
    add_box(b)

# ---- ontology connector edges (zero-thickness in DOM, so synthesized) ----
EDGE = "rgba(255, 255, 255, 0.2)"
n = GEOM["svg"]["nodes"]
xs = sorted(set(round(nd["x"], 1) for nd in n))
W = n[0]["w"]; cy1 = n[0]["y"] + n[0]["h"]/2; bot1 = n[0]["y"] + n[0]["h"]
top2 = n[4]["y"]; cy2 = n[4]["y"] + n[4]["h"]/2
lefts = xs; rights = [lx + W for lx in xs]; cxs = [lx + W/2 for lx in xs]
edges = []
for cy in (cy1, cy2):
    for i in range(3):
        edges.append({"x": rights[i], "y": cy-0.7, "w": lefts[i+1]-rights[i], "h": 1.4, "bg": EDGE, "rx": 0})
for cx in cxs:
    edges.append({"x": cx-0.7, "y": bot1, "w": 1.4, "h": top2-bot1, "bg": EDGE, "rx": 0})
for e in edges:
    add_box(e)

# ---- ontology nodes ----
for nd in GEOM["svg"]["nodes"]:
    add_box({"x": nd["x"], "y": nd["y"], "w": nd["w"], "h": nd["h"], "rx": nd["rx"],
             "bg": "rgba(255, 255, 255, 0.08)", "bc": "rgba(255, 255, 255, 0.34)", "bw": 1.4})

# ---- icons (rasterize SVG -> PNG, embed) ----
FLAME = '<svg xmlns="http://www.w3.org/2000/svg" width="20" height="22" viewBox="0 0 20 22" fill="none"><path d="M10 1C10 1 3 8 3 13.5C3 17.09 6.13 20 10 20C13.87 20 17 17.09 17 13.5C17 8 10 1 10 1Z" fill="#E21836"/><path d="M10 7C10 7 7 10.5 7 13C7 14.66 8.34 16 10 16C11.66 16 13 14.66 13 13C13 10.5 10 7 10 7Z" fill="#FFCA37"/></svg>'
flame_svg = "/tmp/_arch_flame.svg"; open(flame_svg, "w").write(FLAME)
def rasterize(src, out, height):
    subprocess.run(["rsvg-convert", "-h", str(height), src, "-o", out], check=True)
    return out
ICONS = {
    "flame": (flame_svg, 220),
    "databricks-logo-full-color.svg": (os.path.join(HERE, "databricks-logo-full-color.svg"), 160),
    "genie-icon-full-color.svg": (os.path.join(HERE, "genie-icon-full-color.svg"), 200),
    "ai-bi-dashboards-icon-full-color.svg": (os.path.join(HERE, "ai-bi-dashboards-icon-full-color.svg"), 200),
    "unity-catalog-icon-full-color.svg": (os.path.join(HERE, "unity-catalog-icon-full-color.svg"), 200),
    "lakeflow-icon-full-color.svg": (os.path.join(HERE, "lakeflow-icon-full-color.svg"), 200),
}
for img in GEOM["images"]:
    kind = img["kind"]
    if kind not in ICONS:
        continue
    src, hgt = ICONS[kind]
    png = "/tmp/_arch_" + os.path.basename(kind).replace(".svg", "") + ".png"
    rasterize(src, png, hgt)
    slide.shapes.add_picture(png, E(img["x"]), E(img["y"]), E(img["w"]), E(img["h"]))

# ---- text (with the few composite cases handled explicitly) ----
NAVY = "rgb(11, 32, 38)"; BLUE = "rgb(23, 64, 134)"; RED = "rgb(226, 24, 54)"
skip = set()
for i, t in enumerate(GEOM["texts"]):
    tx = t["text"]
    if tx.startswith("One governed platform"):
        skip.add(i)  # replaced by composed title
    if tx == "Superior Plus Propane" or (tx == "AI" and t["fs"] == 18.5):
        skip.add(i)
    if tx.startswith("Powered by"):
        skip.add(i)

# composed title (single line, mixed-color runs)
add_text_runs(24, 49.5, 1050, 23, [
    ("One governed platform, with an ontology specific to ", 18.5, 800, NAVY, "Montserrat", False),
    ("Superior Plus Propane", 18.5, 800, BLUE, "Montserrat", False),
    (" for ", 18.5, 800, NAVY, "Montserrat", False),
    ("AI", 18.5, 800, RED, "Montserrat", False),
], align="start", wrap=False)

# powered-by (two lines, right aligned)
add_text_runs(1199.47, 16, 109.53, 25.5,
              [("POWERED BY\nDATABRICKS DATA + AI", 8.5, 700, "rgb(138, 133, 128)", "Open Sans", False)],
              align="right")

# KPI rows: arrow + label combined so they don't overlap
kpi_labels = [t for t in GEOM["texts"] if abs(t["x"]-1195.5) < 1 and t["fs"] == 10]
kpi_arrows = [t for t in GEOM["texts"] if t["fs"] == 11]
for lab in kpi_labels:
    arr = min(kpi_arrows, key=lambda a: abs(a["y"]-lab["y"]))
    add_text_runs(lab["x"], lab["y"], lab["w"], lab["h"], [
        (arr["text"] + " ", 11, 700, arr["color"], "Open Sans", False),
        (lab["text"], 10, 400, lab["color"], "Open Sans", False),
    ], align="start")
for t in GEOM["texts"]:
    if t in kpi_labels or t in kpi_arrows:
        skip.add(GEOM["texts"].index(t))

for i, t in enumerate(GEOM["texts"]):
    if i in skip:
        continue
    T(t)

# ontology node labels
for lb in GEOM["svg"]["labels"]:
    add_text_runs(lb["x"]-4, lb["y"], lb["w"]+8, lb["h"],
                  [(lb["text"], 9.5, 500, "rgba(255, 255, 255, 0.88)", "Open Sans", False)],
                  align="center")

out = os.path.join(HERE, "superior_arch_cio.pptx")
prs.save(out)
print("saved", out)
